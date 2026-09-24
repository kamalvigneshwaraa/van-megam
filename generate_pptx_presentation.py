#!/usr/bin/env python3
"""
Generate 10-Slide Professional PowerPoint Presentation (PPTX)
Vaan Megam Networks Internship Challenge — TDMA Schedule Planner and Optimizer

Presenter: Kamal Vigneshwaraa S
Department: Information Technology
College: Sri Sairam Institute of Technology, Chennai

Features:
- Clean slide headers with ONLY main headings (removed sub-name lines).
- Flowcharts and step-by-step workflow diagrams on at least 8 slides.
- Zero code snippets, zero URL/file path links, zero image links.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_FILENAME = "TDMA_Schedule_Optimizer_Presentation.pptx"

# Colors
NAVY_PRIMARY = RGBColor(30, 55, 153)       # #1e3799
DARK_BLUE = RGBColor(12, 36, 97)          # #0c2461
TEXT_DARK = RGBColor(44, 62, 80)           # #2c3e50
ACCENT_BLUE = RGBColor(41, 128, 185)       # #2980b9
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(232, 244, 248)          # #e8f4f8
GREEN_BG = RGBColor(234, 250, 234)         # #eafaea
GREEN_BORDER = RGBColor(39, 174, 96)


def add_clean_slide_header(slide, title_text):
    """Adds a clean header banner to content slides with ONLY the title heading."""
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.85))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY_PRIMARY
    top_bar.line.color.rgb = NAVY_PRIMARY

    tf = top_bar.text_frame
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"


def add_horizontal_flowchart(slide, steps, top_y=Inches(1.1), height=Inches(1.2)):
    """Creates a 4 or 5 step horizontal flowchart across the slide."""
    num_steps = len(steps)
    total_width = Inches(9.0)
    start_x = Inches(0.5)
    step_width = (total_width - (Inches(0.2) * (num_steps - 1))) / num_steps

    for i, (title, desc) in enumerate(steps):
        x = start_x + i * (step_width + Inches(0.2))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top_y, step_width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE

        tf = card.text_frame
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        tf.margin_top = Inches(0.1)

        p = tf.paragraphs[0]
        p.text = f"{i+1}. {title}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER


def create_slide_1(prs):
    """Slide 1: Front Title Slide with Student Details."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BLUE
    bg.line.color.rgb = DARK_BLUE

    # Title Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(8.4), Inches(2.3))
    card.fill.solid()
    card.fill.fore_color.rgb = NAVY_PRIMARY
    card.line.color.rgb = ACCENT_BLUE

    tf = card.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p1 = tf.paragraphs[0]
    p1.text = "TDMA Network Schedule Planner & Optimizer"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.font.name = "Arial"
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Wireless Protocol Development Technical Internship Challenge"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(241, 242, 246)
    p2.font.name = "Arial"
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = "Vaan Megam Networks Private Limited (VMN), Chennai"
    p3.font.size = Pt(11)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(200, 214, 229)
    p3.font.name = "Arial"
    p3.alignment = PP_ALIGN.CENTER

    # Student Presenter Info Card
    info_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(3.5), Inches(7.0), Inches(3.4))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = WHITE
    info_card.line.color.rgb = ACCENT_BLUE

    tf_info = info_card.text_frame
    tf_info.margin_left = Inches(0.4)
    tf_info.margin_top = Inches(0.3)

    header = tf_info.paragraphs[0]
    header.text = "STUDENT PRESENTER DETAILS"
    header.font.size = Pt(13)
    header.font.bold = True
    header.font.color.rgb = NAVY_PRIMARY
    header.font.name = "Arial"
    header.alignment = PP_ALIGN.CENTER

    lines = [
        ("Name", "Kamal Vigneshwaraa S"),
        ("Department", "Information Technology"),
        ("College", "Sri Sairam Institute of Technology, Chennai"),
        ("Domain", "Wireless Protocol Development & Graph Optimization")
    ]

    for label, val in lines:
        p = tf_info.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {label}: "
        run1.font.bold = True
        run1.font.size = Pt(12)
        run1.font.color.rgb = DARK_BLUE
        
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(12)
        run2.font.color.rgb = TEXT_DARK


def create_slide_2(prs):
    """Slide 2: Executive Summary & Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "1. Executive Summary & Solution Workflow")

    # Workflow Flowchart at top
    flow_steps = [
        ("Shared Spectrum", "Radio signals collide without scheduling."),
        ("Graph Brain", "Builds physical & conflict graphs."),
        ("DSATUR Engine", "Colors conflict graph to optimize slots."),
        ("EMANE Profile", "Generates XML profiles for emulator.")
    ]
    add_horizontal_flowchart(slide, flow_steps, top_y=Inches(1.0), height=Inches(1.15))

    # Details Cards Below
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = ACCENT_BLUE

    tf = card.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Core Project Deliverables & Approach"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    bullets = [
        "Part 1 (Network Brain): Software engine using graph theory and distance-2 coloring for conflict-free scheduling.",
        "Spatial Reuse Optimization: Maximize simultaneous non-interfering transmissions to minimize frame length.",
        "Part 2 (Physical Emulator Bridge): Convert calculated schedule matrix into native EMANE TDMA Radio Model profiles.",
        "Verification & Containerization: 100% collision-free verification and Docker container deployment."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)


def create_slide_3(prs):
    """Slide 3: TDMA Architecture Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "2. Core Wireless Concepts — TDMA Workflow")

    # TDMA Frame Flowchart
    tdma_steps = [
        ("Frame Init", "Time divided into recurring frames."),
        ("Slot Allocation", "Discrete 1ms slots assigned per node."),
        ("Exclusive Tx", "Node transmits exclusively in slot."),
        ("Clean Reception", "Receivers receive uncorrupted data.")
    ]
    add_horizontal_flowchart(slide, tdma_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = ACCENT_BLUE

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Key Operational Principles"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    points = [
        ("Shared Radio Frequency", "Multiple radio nodes communicate over the exact same frequency channel by sharing time."),
        ("Synchronized Time-Slots", "Discrete, synchronized time-slots prevent simultaneous transmissions by adjacent radios."),
        ("Zero Receiver Collisions", "Correct scheduling guarantees zero signal overlapping and clean packet delivery."),
        ("Clock Synchronization", "All participating radios maintain tight clock sync for precise slot timing boundaries.")
    ]
    for title, desc in points:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = DARK_BLUE

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def create_slide_4(prs):
    """Slide 4: Interference & Collision Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "3. Interference Classification Workflow")

    # Collision Decision Workflow
    col_steps = [
        ("Radio Pair", "Check node pair distance."),
        ("Dist <= 500m", "1-Hop Direct Link (Distance-1)."),
        ("Common Neighbor", "2-Hop Hidden Terminal (Distance-2)."),
        ("Slot Decision", "Assign distinct timeslots!")
    ]
    add_horizontal_flowchart(slide, col_steps, top_y=Inches(1.0), height=Inches(1.15))

    # Two Cards
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(4.3), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = RGBColor(255, 240, 240)
    card1.line.color.rgb = RGBColor(235, 77, 75)

    tf1 = card1.text_frame
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.2)

    p = tf1.paragraphs[0]
    p.text = "Direct Link (Distance-1)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(194, 54, 22)

    d1 = [
        "Radios located directly within range (<= 500m).",
        "1-Hop Neighbors in physical graph.",
        "Transmissions jam each other directly.",
        "MUST be assigned different time-slots."
    ]
    for pt in d1:
        p = tf1.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(2.45), Inches(4.3), Inches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(254, 249, 231)
    card2.line.color.rgb = RGBColor(243, 156, 18)

    tf2 = card2.text_frame
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.2)

    p = tf2.paragraphs[0]
    p.text = "Hidden Terminal (Distance-2)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(211, 84, 0)

    d2 = [
        "Radios transmitting to common receiver (A -> B <- C).",
        "2-Hop Neighbors in physical graph.",
        "Signals collide at intermediate receiver B.",
        "MUST be assigned different time-slots."
    ]
    for pt in d2:
        p = tf2.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)


def create_slide_5(prs):
    """Slide 5: Spatial Reuse Decision Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "4. Spatial Reuse Decision Workflow")

    # Spatial Reuse Flowchart
    reuse_steps = [
        ("Distance Check", "Shortest path distance > 2 hops."),
        ("Interference Check", "Zero signal overlap at any receiver."),
        ("Slot Re-allocation", "Safe timeslot reuse across nodes."),
        ("Capacity Boost", "Dramatically increased network throughput!")
    ]
    add_horizontal_flowchart(slide, reuse_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = GREEN_BG
    box.line.color.rgb = GREEN_BORDER

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Benefits of Spatial Reuse Optimization"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 132, 73)

    points = [
        ("No Receiver Collision", "Radios separated by > 2 hops cause zero signal interference at receivers."),
        ("Frame Length Minimization", "Reduces total unique time-slots required per frame (e.g. 16 nodes in 9 slots)."),
        ("Throughput Multiplication", "Multiple radios transmit simultaneously in parallel across the network."),
        ("NP-Hard Optimization", "Calculates optimal spatial reuse using intelligent graph coloring heuristics.")
    ]
    for title, desc in points:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = RGBColor(20, 90, 50)

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def create_slide_6(prs):
    """Slide 6: System Architecture 5-Stage Flowchart."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "5. System Architecture & 5-Stage Processing Workflow")

    # 5-Stage Flowchart Cards
    stages = [
        ("Stage 1: Input Data", "Receives 2D coordinates for all radio nodes and configures radio range (500m)."),
        ("Stage 2: Graph Build", "Constructs Physical Graph G (1-hop links) & Distance-2 Conflict Graph G_conflict."),
        ("Stage 3: DSATUR Engine", "Executes DSATUR graph coloring algorithm to assign non-conflicting slots."),
        ("Stage 4: Verification", "Performs 100% collision verification check for all 1-hop and 2-hop neighbor pairs."),
        ("Stage 5: EMANE Bridge", "Translates schedule matrix into native EMANE XML configuration profiles.")
    ]

    y_pos = Inches(1.0)
    for i, (title, desc) in enumerate(stages):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), y_pos, Inches(9.0), Inches(1.0))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE

        tf = card.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = f"{i+1}. {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

        y_pos += Inches(1.15)


def create_slide_7(prs):
    """Slide 7: DSATUR Algorithm Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "6. Part 1 — DSATUR Graph Coloring Workflow")

    dsatur_steps = [
        ("Build G_conflict", "Add edge if shortest path is 1 or 2 hops."),
        ("Track Saturation", "Count neighbor assigned colors."),
        ("Pick Max Priority", "Select uncolored node with max saturation."),
        ("Assign Color", "Assign lowest valid integer color index.")
    ]
    add_horizontal_flowchart(slide, dsatur_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = ACCENT_BLUE

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Mathematical Graph Coloring Principles"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    algo = [
        ("Physical Graph G = (V, E)", "Vertices represent radios, edges represent 1-hop physical links (<= 500m)."),
        ("Conflict Graph G_conflict", "Edge added if path is 1 or 2 hops. Nodes > 2 hops share no edge (allows spatial reuse)."),
        ("Saturation Degree Metric", "Dynamic priority metric tracking number of distinct colors assigned to neighbors."),
        ("Lowest Valid Color", "Assigns lowest integer color index {0, 1, 2, ...} not used by neighbors in G_conflict.")
    ]
    for title, desc in algo:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = DARK_BLUE

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def create_slide_8(prs):
    """Slide 8: Verification & Structural Matrix Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "7. Part 1 — Verification & Report Workflow")

    verif_steps = [
        ("Exhaustive Pair Check", "Check all 1-hop & 2-hop pairs."),
        ("Collision Freedom", "Confirm zero shared timeslots."),
        ("Frame Optimization", "Verify minimal frame length."),
        ("Matrix Output", "Generate Slot x Node Boolean matrix.")
    ]
    add_horizontal_flowchart(slide, verif_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = GREEN_BG
    box.line.color.rgb = GREEN_BORDER

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Verification Engine Outcomes"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(20, 90, 50)

    outcomes = [
        ("100% Conflict-Free Verification", "Exhaustively verifies all 1-hop and 2-hop node pairs across topology."),
        ("Zero Receiver Collisions", "Guarantees zero direct link and zero hidden terminal collisions."),
        ("Optimized Frame Length", "Reduces 16-node grid frame length down to 9 unique timeslots."),
        ("Structural Matrix Representation", "Produces formatted Slot x Node Boolean schedule matrix representation."),
        ("Clean Execution Status", "Reports confirmation: Execution finalized cleanly. Schedule verified conflict-free.")
    ]
    for title, desc in outcomes:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = RGBColor(20, 90, 50)

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def create_slide_9(prs):
    """Slide 9: EMANE Emulator Integration Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "8. Part 2 — EMANE Integration Workflow")

    emane_steps = [
        ("Schedule Matrix", "Receive Python calculated slots."),
        ("XML Profile Generator", "Create 1ms slots & NEM tx mapping."),
        ("Multicast Event Bridge", "Publish events to EMANE daemon."),
        ("Docker Container", "Single-command deployment setup.")
    ]
    add_horizontal_flowchart(slide, emane_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = ACCENT_BLUE

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "EMANE Emulation Integration Details"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    emane = [
        ("EMANE Emulation Framework", "Extendable Mobile Ad-hoc Network Emulator for real-time wireless protocol testing."),
        ("Native Profile Generation", "Translates node-to-slot mapping into native XML schedule profiles."),
        ("1ms Slot & NEM Mapping", "Configures 1ms slot duration, 1MHz bandwidth, 2.4GHz frequency, per-NEM slots."),
        ("Multicast Event Publisher", "Broadcasts schedule updates to EMANE event daemons over UDP multicast."),
        ("Docker Containerization", "Consolidated Dockerfile & Docker Compose configuration for containerized deployment.")
    ]
    for title, desc in emane:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = DARK_BLUE

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def create_slide_10(prs):
    """Slide 10: Conclusion & Key Outcomes Workflow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_clean_slide_header(slide, "9. Conclusion & Project Outcomes")

    out_steps = [
        ("Math Optimization", "Graph coloring with spatial reuse."),
        ("Zero Collisions", "Distance-1 & Distance-2 eliminated."),
        ("EMANE Integration", "Native XML profiles generated."),
        ("Real Deployment", "Applicable to 5G & MANET tactical networks.")
    ]
    add_horizontal_flowchart(slide, out_steps, top_y=Inches(1.0), height=Inches(1.15))

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.45), Inches(9.0), Inches(4.5))
    box.fill.solid()
    box.fill.fore_color.rgb = CARD_BG
    box.line.color.rgb = ACCENT_BLUE

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "Summary of Achievements & Impact"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    impacts = [
        ("Complete Solution", "Successfully solved both Part 1 (Network Brain Engine) and Part 2 (EMANE Physical Emulator Bridge)."),
        ("Mathematical Collision Freedom", "Eliminated direct link and hidden terminal collisions completely."),
        ("Optimal Spatial Reuse", "Maximized parallel non-interfering transmissions, reducing total frame length."),
        ("Standardized Output", "Seamless integration with EMANE emulation environment."),
        ("Real-World Application", "Directly applicable to 5G private networks, MANET tactical communications, and SDR mesh networks.")
    ]
    for title, desc in impacts:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"• {title}: "
        run1.font.bold = True
        run1.font.size = Pt(11.5)
        run1.font.color.rgb = DARK_BLUE

        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    create_slide_1(prs)
    create_slide_2(prs)
    create_slide_3(prs)
    create_slide_4(prs)
    create_slide_5(prs)
    create_slide_6(prs)
    create_slide_7(prs)
    create_slide_8(prs)
    create_slide_9(prs)
    create_slide_10(prs)

    output_path = os.path.join(os.path.dirname(__file__), PPTX_FILENAME)
    prs.save(output_path)
    print(f"[+] Successfully generated flowchart PowerPoint presentation: {output_path}")


if __name__ == "__main__":
    build_presentation()
